from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    DARK_BLUE = RGBColor(10, 20, 40)
    GOLD = RGBColor(255, 215, 0)
    LIGHT_BLUE = RGBColor(26, 42, 74)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(180, 180, 180)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_title_slide():
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        title = slide.shapes.title
        title.text = "🏴‍☠️ 海贼王知识库问答系统"
        title.text_frame.paragraphs[0].font.size = Pt(40)
        title.text_frame.paragraphs[0].font.color.rgb = GOLD
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        subtitle = slide.placeholders[1]
        subtitle.text = "基于AI大模型的智能问答助手\n探索伟大航路，解答你的海贼王疑问"
        for p in subtitle.text_frame.paragraphs:
            p.font.size = Pt(20)
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    def add_feature_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "核心功能"
        p.font.size = Pt(32)
        p.font.color.rgb = GOLD
        p.font.bold = True
        
        features = [
            ("🤖 智能问答", "基于大语言模型，理解用户自然语言提问，提供精准、详细的回答"),
            ("📚 丰富知识库", "包含人物、果实、事件、世界观等10+类专业数据"),
            ("🔍 精准检索", "基于向量数据库的语义检索，快速匹配相关知识"),
            ("💬 对话交互", "支持多轮对话、历史记录、快捷提问"),
            ("🚀 流式输出", "支持SSE流式响应，答案实时输出"),
            ("📱 响应式设计", "完美适配桌面端和移动端")
        ]
        
        cols = 3
        rows = 2
        left = Inches(0.8)
        top = Inches(1.4)
        width = Inches((13.33 - 1.6) / cols)
        height = Inches((7.5 - 2) / rows)
        
        for i, (title, desc) in enumerate(features):
            col = i % cols
            row = i // cols
            box_left = left + col * width
            box_top = top + row * height
            
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                          box_left, box_top, width - Inches(0.3), height - Inches(0.3))
            shape.fill.solid()
            shape.fill.fore_color.rgb = LIGHT_BLUE
            shape.line.color.rgb = GOLD
            shape.line.width = Pt(2)
            shape.rounding = Inches(0.2)
            
            text_width = width - Inches(0.6)
            text_height = height - Inches(0.6)
            
            txBox = slide.shapes.add_textbox(box_left + Inches(0.3), box_top + Inches(0.3), 
                                            text_width, text_height)
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p_title = tf.add_paragraph()
            p_title.text = title
            p_title.font.size = Pt(16)
            p_title.font.color.rgb = GOLD
            p_title.font.bold = True
            p_title.space_after = Pt(6)
            
            p_desc = tf.add_paragraph()
            p_desc.text = desc
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = WHITE
            p_desc.word_wrap = True

    def add_tech_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "技术架构"
        p.font.size = Pt(32)
        p.font.color.rgb = GOLD
        p.font.bold = True
        
        techs = [
            ("FastAPI", "高性能Python Web框架，支持异步处理"),
            ("LangChain", "LLM应用开发框架，提供RAG能力"),
            ("ChromaDB", "轻量级向量数据库，高效存储检索"),
            ("DashScope Embeddings", "阿里云向量嵌入服务"),
            ("OpenAI LLM", "大语言模型，强大的自然语言能力"),
            ("HTML/CSS/JS", "现代化前端技术，精美交互界面")
        ]
        
        left = Inches(1)
        top = Inches(1.4)
        width = Inches(11)
        height = Inches(0.9)
        
        for i, (name, desc) in enumerate(techs):
            box_top = top + i * (height + Inches(0.15))
            
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, box_top, Inches(0.3), height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = GOLD
            
            txBox = slide.shapes.add_textbox(left + Inches(0.5), box_top, width - Inches(0.5), height + Inches(0.2))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p_name = tf.add_paragraph()
            p_name.text = name
            p_name.font.size = Pt(16)
            p_name.font.color.rgb = GOLD
            p_name.font.bold = True
            
            p_desc = tf.add_paragraph()
            p_desc.text = desc
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = WHITE
            p_desc.word_wrap = True

    def add_data_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "知识库内容"
        p.font.size = Pt(32)
        p.font.color.rgb = GOLD
        p.font.bold = True
        
        p_sub = tf.add_paragraph()
        p_sub.text = "10大类专业数据，涵盖海贼王全系列"
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = GRAY
        
        data_items = [
            ("👤 人物库", "主要角色详细档案"),
            ("🍎 恶魔果实", "果实能力详解"),
            ("👑 四皇档案", "四皇势力完整资料"),
            ("⚓ 七武海", "历代王下七武海详解"),
            ("🔱 海军高层", "海军本部将领档案"),
            ("📢 革命军", "革命军势力完整档案"),
            ("🚢 海贼团", "各大海贼团介绍"),
            ("🌑 黑暗世界", "地下势力帝王档案"),
            ("📅 重大事件", "时间线完整梳理"),
            ("🌍 世界观", "海贼王世界框架介绍")
        ]
        
        cols = 2
        rows = 5
        left = Inches(1)
        top = Inches(1.8)
        width = Inches((13.33 - 2) / cols)
        height = Inches((7.5 - 2.2) / rows)
        
        for i, (icon_name, desc) in enumerate(data_items):
            col = i % cols
            row = i // cols
            box_left = left + col * width
            box_top = top + row * height
            
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                          box_left, box_top, width - Inches(0.3), height - Inches(0.2))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(20, 35, 65)
            shape.line.color.rgb = RGBColor(50, 65, 95)
            shape.rounding = Inches(0.15)
            
            txBox = slide.shapes.add_textbox(box_left + Inches(0.3), box_top + Inches(0.15), 
                                            width - Inches(0.9), height - Inches(0.5))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p_icon = tf.add_paragraph()
            p_icon.text = icon_name
            p_icon.font.size = Pt(14)
            p_icon.font.color.rgb = WHITE
            p_icon.font.bold = True
            
            p_desc = tf.add_paragraph()
            p_desc.text = desc
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = GRAY
            p_desc.word_wrap = True

    def add_ui_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "界面展示"
        p.font.size = Pt(32)
        p.font.color.rgb = GOLD
        p.font.bold = True
        
        p_sub = tf.add_paragraph()
        p_sub.text = "精美设计，海贼王主题风格"
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = GRAY
        
        ui_features = [
            "深蓝色航海主题背景，金色文字点缀",
            "知识导航、人物卡片、历史记录三大侧边栏模块",
            "热门问题快捷入口，一键提问",
            "流畅的对话交互体验，支持流式输出",
            "响应式布局，适配各种屏幕尺寸"
        ]
        
        left = Inches(1)
        top = Inches(1.4)
        width = Inches(5.5)
        height = Inches(0.7)
        
        for i, feature in enumerate(ui_features):
            box_top = top + i * (height + Inches(0.1))
            
            txBox = slide.shapes.add_textbox(left, box_top, width, height + Inches(0.2))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p_feature = tf.add_paragraph()
            p_feature.text = "✦ " + feature
            p_feature.font.size = Pt(14)
            p_feature.font.color.rgb = WHITE
            p_feature.word_wrap = True
        
        preview_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                            Inches(7.5), Inches(1.4), Inches(5.5), Inches(5.3))
        preview_box.fill.solid()
        preview_box.fill.fore_color.rgb = LIGHT_BLUE
        preview_box.line.color.rgb = GOLD
        preview_box.line.width = Pt(2)
        preview_box.rounding = Inches(0.3)
        
        txBox = slide.shapes.add_textbox(Inches(7.8), Inches(1.9), Inches(5), Inches(4))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p_title = tf.add_paragraph()
        p_title.text = "🏴‍☠️ 海贼王知识库"
        p_title.font.size = Pt(18)
        p_title.font.color.rgb = GOLD
        p_title.font.bold = True
        p_title.alignment = PP_ALIGN.CENTER
        p_title.space_after = Pt(12)
        
        p_chat = tf.add_paragraph()
        p_chat.text = "你好！我是海贼王专属知识库问答机器人..."
        p_chat.font.size = Pt(12)
        p_chat.font.color.rgb = WHITE
        p_chat.space_after = Pt(8)
        p_chat.word_wrap = True
        
        input_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 
                                          Inches(7.8), Inches(6.2), Inches(4), Inches(0.6))
        input_box.fill.solid()
        input_box.fill.fore_color.rgb = RGBColor(40, 55, 85)
        input_box.line.color.rgb = GOLD
        
        send_btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                         Inches(12), Inches(6.2), Inches(0.8), Inches(0.6))
        send_btn.fill.solid()
        send_btn.fill.fore_color.rgb = RGBColor(255, 215, 50)
        send_btn.rounding = Inches(0.1)

    def add_use_case_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "使用场景"
        p.font.size = Pt(32)
        p.font.color.rgb = GOLD
        p.font.bold = True
        
        p_sub = tf.add_paragraph()
        p_sub.text = "适用于多种海贼王知识查询需求"
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = GRAY
        
        use_cases = [
            ("🎮 动漫迷查询", "快速了解人物背景、能力设定、剧情走向"),
            ("📝 创作参考", "为同人创作提供准确的人物设定和世界观资料"),
            ("👨‍🏫 知识学习", "系统化学习海贼王知识体系，深入理解作品内涵"),
            ("👥 社群讨论", "为粉丝社群讨论提供准确的数据支撑")
        ]
        
        cols = 2
        rows = 2
        left = Inches(1)
        top = Inches(1.4)
        width = Inches((13.33 - 2) / cols)
        height = Inches((7.5 - 2) / rows)
        
        for i, (title, desc) in enumerate(use_cases):
            col = i % cols
            row = i // cols
            box_left = left + col * width
            box_top = top + row * height
            
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                          box_left, box_top, width - Inches(0.3), height - Inches(0.3))
            shape.fill.solid()
            shape.fill.fore_color.rgb = LIGHT_BLUE
            shape.line.color.rgb = GOLD
            shape.rounding = Inches(0.3)
            
            txBox = slide.shapes.add_textbox(box_left + Inches(0.4), box_top + Inches(0.3), 
                                            width - Inches(1), height - Inches(0.8))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p_title = tf.add_paragraph()
            p_title.text = title
            p_title.font.size = Pt(18)
            p_title.font.color.rgb = GOLD
            p_title.font.bold = True
            p_title.alignment = PP_ALIGN.CENTER
            p_title.space_after = Pt(8)
            
            p_desc = tf.add_paragraph()
            p_desc.text = desc
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = WHITE
            p_desc.alignment = PP_ALIGN.CENTER
            p_desc.word_wrap = True

    def add_summary_slide():
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        set_slide_background(slide, DARK_BLUE)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p_emoji = tf.add_paragraph()
        p_emoji.text = "🎉"
        p_emoji.font.size = Pt(50)
        p_emoji.alignment = PP_ALIGN.CENTER
        p_emoji.space_after = Pt(8)
        
        p_title = tf.add_paragraph()
        p_title.text = "探索伟大航路"
        p_title.font.size = Pt(32)
        p_title.font.color.rgb = GOLD
        p_title.font.bold = True
        p_title.alignment = PP_ALIGN.CENTER
        
        txBox2 = slide.shapes.add_textbox(Inches(2), Inches(2.8), Inches(9), Inches(2.5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        
        p_content = tf2.add_paragraph()
        p_content.text = ("海贼王知识库问答系统，为所有海贼迷提供专业、准确、全面的海贼王知识查询服务。\n"
                         "无论是人物档案、恶魔果实、重大事件还是世界观设定，都能为您一一解答！")
        p_content.font.size = Pt(16)
        p_content.font.color.rgb = WHITE
        p_content.alignment = PP_ALIGN.CENTER
        p_content.line_spacing = Pt(16)
        p_content.word_wrap = True
        
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                      Inches(4.6), Inches(5.8), Inches(4), Inches(0.8))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(40, 60, 100)
        badge.line.color.rgb = GOLD
        badge.line.width = Pt(2)
        badge.rounding = Inches(0.4)
        
        txBox3 = slide.shapes.add_textbox(Inches(4.8), Inches(5.9), Inches(3.6), Inches(0.6))
        tf3 = txBox3.text_frame
        
        p_badge = tf3.add_paragraph()
        p_badge.text = "ONE PIECE!"
        p_badge.font.size = Pt(22)
        p_badge.font.color.rgb = GOLD
        p_badge.font.bold = True
        p_badge.alignment = PP_ALIGN.CENTER

    add_title_slide()
    add_feature_slide()
    add_tech_slide()
    add_data_slide()
    add_ui_slide()
    add_use_case_slide()
    add_summary_slide()

    output_path = "海贼王知识库问答系统_功能介绍.pptx"
    prs.save(output_path)
    print(f"PPT文件已生成：{output_path}")

if __name__ == "__main__":
    create_ppt()